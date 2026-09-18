"""كاتب DOCX عربي احترافي: فقرات RTL حقيقية (<w:bidi/> و<w:rtl/>) وعناوين وقوائم.

هذه هي الوحدة التي تحل مشكلة «الوورد العربي المضروب»:
- كل فقرة عربية تُنشأ بـ w:bidi (اتجاه فقرة RTL في Word) ومحاذاة يمنى.
- كل تشغيلة (run) تحمل w:rtl + خط Complex Script بحجمها → تُعرض الحروف متصلة صحيحة.
- الأسطر المختلطة تبقى سليمة: Word يطبق خوارزمية Unicode BiDi على النص المنطقي.
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from ..core.languages import dominant_script

DEFAULT_AR_FONT = "Traditional Arabic"
DEFAULT_AR_SIZE = Pt(14)


def _set_rtl_paragraph(p, rtl: bool):
    pPr = p._p.get_or_add_pPr()
    for el in pPr.findall(qn("w:bidi")):
        pPr.remove(el)
    if rtl:
        pPr.append(pPr.makeelement(qn("w:bidi"), {}))


def _set_rtl_run(run, rtl: bool, ar_font: str | None):
    rPr = run._r.get_or_add_rPr()
    for el in rPr.findall(qn("w:rtl")):
        rPr.remove(el)
    if rtl:
        rPr.append(rPr.makeelement(qn("w:rtl"), {}))
        if ar_font:
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = rPr.makeelement(qn("w:rFonts"), {})
                rPr.insert(0, rFonts)
            rFonts.set(qn("w:cs"), ar_font)


def _set_run_size(run, size):
    run.font.size = size
    rPr = run._r.get_or_add_rPr()
    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = rPr.makeelement(qn("w:szCs"), {})
        rPr.append(szCs)
    szCs.set(qn("w:val"), str(int(size.pt * 2)))


class DocxWriter:
    def __init__(self, path: str, ar_font: str = DEFAULT_AR_FONT, ar_size=DEFAULT_AR_SIZE,
                 page_breaks_between_pages: bool = False):
        self.path = path
        self.doc = Document()
        self.ar_font = ar_font
        self.ar_size = ar_size
        self.page_breaks = page_breaks_between_pages
        self._last_page = None
        self._list_open = False

    def _maybe_page_break(self, page):
        if self.page_breaks and self._last_page is not None and page is not None and page > self._last_page:
            self.doc.add_page_break()
        if page is not None:
            self._last_page = page

    def add_block(self, block: dict):
        btype = block["type"]
        text = block["text"]
        page = block.get("page")
        self._maybe_page_break(page)
        rtl = dominant_script(text) == "ar"
        if btype == "heading":
            p = self.doc.add_paragraph()
            _set_rtl_paragraph(p, rtl)
            if rtl:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(text)
            run.bold = True
            _set_rtl_run(run, rtl, self.ar_font)
            _set_run_size(run, self.ar_size)
        elif btype == "list_item":
            p = self.doc.add_paragraph(style="List Bullet" if not rtl else None)
            _set_rtl_paragraph(p, rtl)
            if rtl:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run("• " + text if rtl else text)
            _set_rtl_run(run, rtl, self.ar_font)
            _set_run_size(run, self.ar_size)
        else:
            p = self.doc.add_paragraph()
            _set_rtl_paragraph(p, rtl)
            if rtl:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(text)
            _set_rtl_run(run, rtl, self.ar_font)
            _set_run_size(run, self.ar_size)

    def add_blocks(self, blocks):
        for b in blocks:
            self.add_block(b)

    def save(self):
        self.doc.save(self.path)
